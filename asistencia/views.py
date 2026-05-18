import json
import io
from datetime import datetime, timedelta
from django.http import JsonResponse, HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.views.decorators.csrf import csrf_exempt
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from .models import Empleado, Registro


# ============================================================
# VISTAS PÚBLICAS
# ============================================================

def pantalla_principal(request):
    """Pantalla del checador con registros del día."""
    hoy = timezone.localdate()
    registros_hoy = (
        Registro.objects
        .filter(fecha_hora__date=hoy)
        .select_related('empleado')
        .order_by('-fecha_hora')
    )
    contexto = {
        'registros_hoy': registros_hoy,
        'total_hoy': registros_hoy.count(),
        'fecha_hoy': hoy,
    }
    return render(request, 'index.html', contexto)


def obtener_registros_hoy(request):
    """API JSON con los registros del día (para auto-actualizar)."""
    hoy = timezone.localdate()
    registros = (
        Registro.objects
        .filter(fecha_hora__date=hoy)
        .select_related('empleado')
        .order_by('-fecha_hora')
    )
    data = [
        {
            'empleado': r.empleado.nombre,
            'tipo': r.tipo,
            'hora': timezone.localtime(r.fecha_hora).strftime('%H:%M:%S'),
        }
        for r in registros
    ]
    return JsonResponse({'registros': data, 'total': len(data)})


@csrf_exempt
def registrar_chequeo(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            uid_recibido = data.get('uid')
            empleado = Empleado.objects.filter(uid_rfid=uid_recibido).first()

            if not empleado:
                return JsonResponse({'status': 'error', 'mensaje': 'No registrado'}, status=404)

            ultimo = Registro.objects.filter(empleado=empleado).order_by('-fecha_hora').first()
            nuevo_tipo = 'SALIDA' if ultimo and ultimo.tipo == 'ENTRADA' else 'ENTRADA'

            Registro.objects.create(empleado=empleado, tipo=nuevo_tipo)
            return JsonResponse({'status': 'success', 'empleado': empleado.nombre, 'accion': nuevo_tipo})
        except Exception as e:
            return JsonResponse({'status': 'error', 'mensaje': str(e)}, status=400)
    return JsonResponse({'status': 'error', 'mensaje': 'POST requerido'}, status=405)


# ============================================================
# SISTEMA DEL PATRÓN — LOGIN / LOGOUT
# ============================================================

def patron_login(request):
    """Página de login para el patrón."""
    if request.user.is_authenticated:
        return redirect('patron_dashboard')

    error = None
    if request.method == 'POST':
        usuario = request.POST.get('usuario')
        contrasena = request.POST.get('contrasena')
        user = authenticate(request, username=usuario, password=contrasena)
        if user is not None and user.is_staff:
            login(request, user)
            return redirect('patron_dashboard')
        else:
            error = 'Usuario o contraseña incorrectos.'

    return render(request, 'patron/login.html', {'error': error})


@login_required(login_url='patron_login')
def patron_logout(request):
    """Cierra la sesión del patrón."""
    logout(request)
    return redirect('patron_login')


# ============================================================
# DASHBOARD DEL PATRÓN
# ============================================================

@login_required(login_url='patron_login')
def patron_dashboard(request):
    """Página principal del patrón con lista de barberos y registros."""
    if not request.user.is_staff:
        return redirect('patron_login')

    barberos = Empleado.objects.all().order_by('nombre')
    hoy = timezone.localdate()

    # Filtros de registros
    fecha_filtro = request.GET.get('fecha', hoy.isoformat())
    barbero_filtro = request.GET.get('barbero', 'todos')

    registros = Registro.objects.select_related('empleado').order_by('-fecha_hora')
    if fecha_filtro:
        registros = registros.filter(fecha_hora__date=fecha_filtro)
    if barbero_filtro and barbero_filtro != 'todos':
        registros = registros.filter(empleado_id=barbero_filtro)

    contexto = {
        'barberos': barberos,
        'registros': registros[:100],  # límite de 100 por rendimiento
        'fecha_filtro': fecha_filtro,
        'barbero_filtro': barbero_filtro,
        'total_barberos': barberos.count(),
        'total_registros_hoy': Registro.objects.filter(fecha_hora__date=hoy).count(),
    }
    return render(request, 'patron/dashboard.html', contexto)


# ============================================================
# CRUD DE BARBEROS
# ============================================================

@login_required(login_url='patron_login')
def patron_agregar_barbero(request):
    """Agregar un nuevo barbero."""
    if not request.user.is_staff:
        return redirect('patron_login')

    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        uid_rfid = request.POST.get('uid_rfid', '').strip().upper()

        if not nombre or not uid_rfid:
            messages.error(request, 'Todos los campos son obligatorios.')
            return redirect('patron_agregar_barbero')

        if Empleado.objects.filter(uid_rfid=uid_rfid).exists():
            messages.error(request, f'Ya existe un barbero con el UID {uid_rfid}.')
            return redirect('patron_agregar_barbero')

        Empleado.objects.create(nombre=nombre, uid_rfid=uid_rfid)
        messages.success(request, f'✂ Barbero "{nombre}" agregado correctamente.')
        return redirect('patron_dashboard')

    return render(request, 'patron/agregar_barbero.html')


@login_required(login_url='patron_login')
def patron_editar_barbero(request, barbero_id):
    """Editar un barbero existente."""
    if not request.user.is_staff:
        return redirect('patron_login')

    barbero = get_object_or_404(Empleado, id=barbero_id)

    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        uid_rfid = request.POST.get('uid_rfid', '').strip().upper()

        if not nombre or not uid_rfid:
            messages.error(request, 'Todos los campos son obligatorios.')
            return redirect('patron_editar_barbero', barbero_id=barbero_id)

        # Validar UID único (excluyendo al barbero actual)
        if Empleado.objects.filter(uid_rfid=uid_rfid).exclude(id=barbero_id).exists():
            messages.error(request, f'Ya existe otro barbero con el UID {uid_rfid}.')
            return redirect('patron_editar_barbero', barbero_id=barbero_id)

        barbero.nombre = nombre
        barbero.uid_rfid = uid_rfid
        barbero.save()
        messages.success(request, f'✂ Barbero "{nombre}" actualizado correctamente.')
        return redirect('patron_dashboard')

    return render(request, 'patron/editar_barbero.html', {'barbero': barbero})


@login_required(login_url='patron_login')
def patron_eliminar_barbero(request, barbero_id):
    """Eliminar un barbero y todos sus registros."""
    if not request.user.is_staff:
        return redirect('patron_login')

    barbero = get_object_or_404(Empleado, id=barbero_id)

    if request.method == 'POST':
        nombre = barbero.nombre
        total_registros = barbero.registro_set.count()
        barbero.delete()  # CASCADE borra también sus registros
        messages.success(
            request,
            f'🗑 Barbero "{nombre}" eliminado junto con {total_registros} registro(s).'
        )
        return redirect('patron_dashboard')

    return render(request, 'patron/eliminar_barbero.html', {'barbero': barbero})


# ============================================================
# CRUD DE REGISTROS DE ASISTENCIA
# ============================================================

@login_required(login_url='patron_login')
def patron_editar_registro(request, registro_id):
    """Editar la fecha/hora o tipo de un registro."""
    if not request.user.is_staff:
        return redirect('patron_login')

    registro = get_object_or_404(Registro, id=registro_id)

    if request.method == 'POST':
        fecha_str = request.POST.get('fecha')
        hora_str = request.POST.get('hora')
        tipo = request.POST.get('tipo')

        try:
            fecha_hora_str = f"{fecha_str} {hora_str}"
            nueva_fecha_hora = datetime.strptime(fecha_hora_str, '%Y-%m-%d %H:%M')
            # Convertir a timezone-aware
            nueva_fecha_hora = timezone.make_aware(nueva_fecha_hora)

            registro.fecha_hora = nueva_fecha_hora
            registro.tipo = tipo
            registro.save()

            messages.success(request, '✓ Registro actualizado correctamente.')
            return redirect('patron_dashboard')
        except (ValueError, TypeError):
            messages.error(request, 'Formato de fecha u hora inválido.')

    # Hora local para el formulario
    fecha_hora_local = timezone.localtime(registro.fecha_hora)
    contexto = {
        'registro': registro,
        'fecha_actual': fecha_hora_local.strftime('%Y-%m-%d'),
        'hora_actual': fecha_hora_local.strftime('%H:%M'),
    }
    return render(request, 'patron/editar_registro.html', contexto)


@login_required(login_url='patron_login')
def patron_eliminar_registro(request, registro_id):
    """Eliminar un registro de asistencia."""
    if not request.user.is_staff:
        return redirect('patron_login')

    registro = get_object_or_404(Registro, id=registro_id)

    if request.method == 'POST':
        registro.delete()
        messages.success(request, '🗑 Registro eliminado correctamente.')
        return redirect('patron_dashboard')

    return render(request, 'patron/eliminar_registro.html', {'registro': registro})


@login_required(login_url='patron_login')
def patron_crear_registro_manual(request):
    """Crear un registro manual (cuando un barbero olvida checar)."""
    if not request.user.is_staff:
        return redirect('patron_login')

    barberos = Empleado.objects.all().order_by('nombre')

    if request.method == 'POST':
        barbero_id = request.POST.get('barbero_id')
        fecha_str = request.POST.get('fecha')
        hora_str = request.POST.get('hora')
        tipo = request.POST.get('tipo')

        try:
            barbero = Empleado.objects.get(id=barbero_id)
            fecha_hora_str = f"{fecha_str} {hora_str}"
            fecha_hora = datetime.strptime(fecha_hora_str, '%Y-%m-%d %H:%M')
            fecha_hora = timezone.make_aware(fecha_hora)

            Registro.objects.create(
                empleado=barbero,
                fecha_hora=fecha_hora,
                tipo=tipo,
            )
            messages.success(
                request,
                f'✓ Registro manual creado para {barbero.nombre} ({tipo}).'
            )
            return redirect('patron_dashboard')
        except (Empleado.DoesNotExist, ValueError, TypeError):
            messages.error(request, 'Datos inválidos. Revisa el formulario.')

    hoy = timezone.localdate()
    ahora = timezone.localtime()
    contexto = {
        'barberos': barberos,
        'fecha_default': hoy.isoformat(),
        'hora_default': ahora.strftime('%H:%M'),
    }
    return render(request, 'patron/crear_registro.html', contexto)


# ============================================================
# REPORTES (PDF / EXCEL)
# ============================================================

def _obtener_registros_filtrados(request):
    """Aplica filtros de fecha y empleado a los registros."""
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')
    empleado_id = request.GET.get('empleado_id')

    registros = Registro.objects.select_related('empleado').order_by('fecha_hora')

    if fecha_inicio:
        registros = registros.filter(fecha_hora__date__gte=fecha_inicio)
    if fecha_fin:
        registros = registros.filter(fecha_hora__date__lte=fecha_fin)
    if empleado_id and empleado_id != 'todos':
        registros = registros.filter(empleado_id=empleado_id)

    return registros, fecha_inicio, fecha_fin, empleado_id


def _calcular_horas_trabajadas(registros):
    """Empareja entradas con salidas para calcular horas trabajadas."""
    filas = []
    pendientes_entrada = {}

    for r in registros:
        emp_id = r.empleado_id
        hora_local = timezone.localtime(r.fecha_hora)

        if r.tipo == 'ENTRADA':
            pendientes_entrada[emp_id] = r
            filas.append({
                'empleado': r.empleado.nombre,
                'fecha': hora_local.strftime('%d/%m/%Y'),
                'tipo': 'Entrada',
                'hora': hora_local.strftime('%H:%M:%S'),
                'horas_trabajadas': '',
            })
        else:
            horas_str = ''
            if emp_id in pendientes_entrada:
                entrada = pendientes_entrada.pop(emp_id)
                delta = r.fecha_hora - entrada.fecha_hora
                horas = delta.total_seconds() / 3600
                horas_str = f"{horas:.2f} h"

            filas.append({
                'empleado': r.empleado.nombre,
                'fecha': hora_local.strftime('%d/%m/%Y'),
                'tipo': 'Salida',
                'hora': hora_local.strftime('%H:%M:%S'),
                'horas_trabajadas': horas_str,
            })

    return filas


@login_required(login_url='patron_login')
def pagina_reportes(request):
    """Página de generación de reportes."""
    if not request.user.is_staff:
        return redirect('patron_login')

    empleados = Empleado.objects.all().order_by('nombre')
    hoy = timezone.localdate()
    primer_dia_mes = hoy.replace(day=1)

    contexto = {
        'empleados': empleados,
        'fecha_inicio_default': primer_dia_mes.isoformat(),
        'fecha_fin_default': hoy.isoformat(),
    }
    return render(request, 'patron/reportes.html', contexto)


@login_required(login_url='patron_login')
def reporte_excel(request):
    """Genera el reporte en Excel."""
    if not request.user.is_staff:
        return redirect('patron_login')

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    registros, fecha_inicio, fecha_fin, empleado_id = _obtener_registros_filtrados(request)
    filas = _calcular_horas_trabajadas(registros)

    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencias"

    encabezado_font = Font(name='Calibri', size=12, bold=True, color='F4E4C1')
    encabezado_fill = PatternFill(start_color='2B1810', end_color='2B1810', fill_type='solid')
    titulo_font = Font(name='Calibri', size=18, bold=True, color='2B1810')
    subtitulo_font = Font(name='Calibri', size=11, italic=True, color='6B4423')
    centrado = Alignment(horizontal='center', vertical='center')
    borde = Border(
        left=Side(style='thin', color='8B6F3A'),
        right=Side(style='thin', color='8B6F3A'),
        top=Side(style='thin', color='8B6F3A'),
        bottom=Side(style='thin', color='8B6F3A'),
    )

    ws.merge_cells('A1:E1')
    ws['A1'] = 'JOHNNY BOY BARBER SHOP'
    ws['A1'].font = titulo_font
    ws['A1'].alignment = centrado

    ws.merge_cells('A2:E2')
    periodo_txt = f"Reporte de Asistencias  |  Del {fecha_inicio or '...'} al {fecha_fin or '...'}"
    ws['A2'] = periodo_txt
    ws['A2'].font = subtitulo_font
    ws['A2'].alignment = centrado

    ws.append([])

    headers = ['Empleado', 'Fecha', 'Tipo', 'Hora', 'Horas Trabajadas']
    ws.append(headers)
    for col_num, _ in enumerate(headers, 1):
        celda = ws.cell(row=4, column=col_num)
        celda.font = encabezado_font
        celda.fill = encabezado_fill
        celda.alignment = centrado
        celda.border = borde

    fila_excel = 5
    for f in filas:
        ws.cell(row=fila_excel, column=1, value=f['empleado']).border = borde
        ws.cell(row=fila_excel, column=2, value=f['fecha']).border = borde
        ws.cell(row=fila_excel, column=3, value=f['tipo']).border = borde
        ws.cell(row=fila_excel, column=4, value=f['hora']).border = borde
        ws.cell(row=fila_excel, column=5, value=f['horas_trabajadas']).border = borde

        color_fondo = 'E8F5E9' if f['tipo'] == 'Entrada' else 'FFEBEE'
        for col in range(1, 6):
            ws.cell(row=fila_excel, column=col).fill = PatternFill(
                start_color=color_fondo, end_color=color_fondo, fill_type='solid'
            )
            ws.cell(row=fila_excel, column=col).alignment = Alignment(horizontal='center')

        fila_excel += 1

    anchos = {'A': 25, 'B': 14, 'C': 12, 'D': 12, 'E': 18}
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[4].height = 22

    fila_excel += 2
    ws.cell(row=fila_excel, column=1, value='Total de registros:').font = Font(bold=True)
    ws.cell(row=fila_excel, column=2, value=len(filas))
    fila_excel += 1
    ws.cell(row=fila_excel, column=1, value='Generado:').font = Font(bold=True)
    ws.cell(row=fila_excel, column=2, value=timezone.localtime().strftime('%d/%m/%Y %H:%M'))

    nombre_archivo = f"asistencias_{fecha_inicio or 'inicio'}_a_{fecha_fin or 'fin'}.xlsx"
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    wb.save(response)
    return response


@login_required(login_url='patron_login')
def reporte_pdf(request):
    """Genera el reporte en PDF."""
    if not request.user.is_staff:
        return redirect('patron_login')

    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    registros, fecha_inicio, fecha_fin, empleado_id = _obtener_registros_filtrados(request)
    filas = _calcular_horas_trabajadas(registros)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=letter,
        rightMargin=1.5*cm, leftMargin=1.5*cm,
        topMargin=1.5*cm, bottomMargin=1.5*cm
    )

    estilos = getSampleStyleSheet()
    titulo_estilo = ParagraphStyle(
        'TituloBarber', parent=estilos['Title'], fontSize=22,
        textColor=colors.HexColor('#2B1810'), alignment=1,
        spaceAfter=6, fontName='Helvetica-Bold',
    )
    subtitulo_estilo = ParagraphStyle(
        'Sub', parent=estilos['Normal'], fontSize=11,
        textColor=colors.HexColor('#6B4423'), alignment=1,
        spaceAfter=4, fontName='Helvetica-Oblique',
    )
    pie_estilo = ParagraphStyle(
        'Pie', parent=estilos['Normal'], fontSize=9,
        textColor=colors.HexColor('#6B4423'), alignment=1,
        fontName='Helvetica-Oblique',
    )

    elementos = []
    elementos.append(Paragraph("✂  JOHNNY BOY BARBER SHOP  ✂", titulo_estilo))
    elementos.append(Paragraph("Reporte de Asistencias", subtitulo_estilo))
    periodo_txt = f"Del {fecha_inicio or '...'} al {fecha_fin or '...'}"
    elementos.append(Paragraph(periodo_txt, subtitulo_estilo))
    elementos.append(Spacer(1, 0.5*cm))

    encabezados = ['Empleado', 'Fecha', 'Tipo', 'Hora', 'Horas Trab.']
    datos_tabla = [encabezados]
    for f in filas:
        datos_tabla.append([f['empleado'], f['fecha'], f['tipo'], f['hora'], f['horas_trabajadas']])

    if len(datos_tabla) == 1:
        datos_tabla.append(['Sin registros en este periodo', '', '', '', ''])

    tabla = Table(datos_tabla, colWidths=[5*cm, 3*cm, 2.5*cm, 2.5*cm, 3*cm], repeatRows=1)

    estilo_tabla = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2B1810')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#F4E4C1')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, 0), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('ALIGN', (0, 1), (-1, -1), 'CENTER'),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('VALIGN', (0, 1), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 1), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 5),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#8B6F3A')),
        ('BOX', (0, 0), (-1, -1), 1.2, colors.HexColor('#2B1810')),
    ])

    for i, f in enumerate(filas, start=1):
        if f['tipo'] == 'Entrada':
            estilo_tabla.add('BACKGROUND', (0, i), (-1, i), colors.HexColor('#E8F5E9'))
        else:
            estilo_tabla.add('BACKGROUND', (0, i), (-1, i), colors.HexColor('#FFEBEE'))

    tabla.setStyle(estilo_tabla)
    elementos.append(tabla)

    elementos.append(Spacer(1, 0.8*cm))
    elementos.append(Paragraph(
        f"Total de registros: {len(filas)}  |  Generado el {timezone.localtime().strftime('%d/%m/%Y %H:%M')}",
        pie_estilo
    ))

    doc.build(elementos)
    pdf_bytes = buffer.getvalue()
    buffer.close()

    nombre_archivo = f"asistencias_{fecha_inicio or 'inicio'}_a_{fecha_fin or 'fin'}.pdf"
    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    return response